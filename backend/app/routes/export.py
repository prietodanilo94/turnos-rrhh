"""
Export API — Excel and JSON formats for n8n / external tools.

GET /api/export/json?branch_id=X&week_start=Y&week_end=Z
GET /api/export/excel?branch_id=X&week_start=Y&week_end=Z
"""
from datetime import date, timedelta
from fastapi import APIRouter, Depends, Query, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session, joinedload
from typing import Optional
import io
import json

from app.database import get_db
from app.models import Branch, Worker, Schedule, ShiftTemplate, User
from app.auth import require_role

router = APIRouter(prefix="/api/export", tags=["Export"])


def _get_schedules_data(db: Session, branch_id: int, week_start: date, week_end: date) -> dict:
    """Build structured schedule data for export."""
    branch = db.query(Branch).filter(Branch.id == branch_id).first()
    if not branch:
        raise HTTPException(status_code=404, detail="Sucursal no encontrada")

    workers = db.query(Worker).filter(
        Worker.branch_id == branch_id,
        Worker.is_active == True,
    ).order_by(Worker.last_name).all()

    data = {
        "branch_id": branch.id,
        "branch_name": branch.name,
        "branch_code": branch.code,
        "week_start": str(week_start),
        "week_end": str(week_end),
        "exported_at": date.today().isoformat(),
        "workers": [],
    }

    for worker in workers:
        schedules = (
            db.query(Schedule)
            .options(joinedload(Schedule.shift_template))
            .filter(
                Schedule.worker_id == worker.id,
                Schedule.date >= week_start,
                Schedule.date <= week_end,
            )
            .order_by(Schedule.date)
            .all()
        )

        days = {}
        total_hours = 0.0
        for s in schedules:
            hours = float(s.total_hours) if s.total_hours else 0.0
            days[str(s.date)] = {
                "is_day_off": s.is_day_off,
                "is_locked": s.is_locked,
                "shift_code": s.shift_template.code if s.shift_template else None,
                "shift_name": s.shift_template.name if s.shift_template else None,
                "start_time": str(s.custom_start_time or (s.shift_template.start_time if s.shift_template else None)),
                "end_time": str(s.custom_end_time or (s.shift_template.end_time if s.shift_template else None)),
                "total_hours": hours,
                "status": s.status,
                "notes": s.notes,
            }
            if not s.is_day_off:
                total_hours += hours

        data["workers"].append({
            "worker_id": worker.id,
            "rut": worker.rut,
            "name": worker.full_name,
            "position": worker.position,
            "contract_type": worker.contract_type,
            "contracted_weekly_hours": worker.contracted_weekly_hours,
            "total_hours_scheduled": round(total_hours, 2),
            "days": days,
        })

    return data


@router.get("/json")
def export_json(
    branch_id: int = Query(...),
    week_start: date = Query(...),
    week_end: Optional[date] = Query(None),
    db: Session = Depends(get_db),
    user: User = Depends(require_role("admin", "manager")),
):
    """
    JSON export designed for n8n / AI agent consumption.
    Includes full worker schedule with shift details, hours, status.
    """
    if week_end is None:
        week_end = week_start + timedelta(days=6)

    data = _get_schedules_data(db, branch_id, week_start, week_end)
    return data


@router.get("/excel")
def export_excel(
    branch_id: int = Query(...),
    week_start: date = Query(...),
    week_end: Optional[date] = Query(None),
    db: Session = Depends(get_db),
    user: User = Depends(require_role("admin", "manager")),
):
    """Download Excel file with the weekly schedule."""
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    if week_end is None:
        week_end = week_start + timedelta(days=6)

    data = _get_schedules_data(db, branch_id, week_start, week_end)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Semana {week_start}"

    # Styles
    header_fill = PatternFill("solid", fgColor="1E293B")
    header_font = Font(color="F1F5F9", bold=True)
    day_off_fill = PatternFill("solid", fgColor="334155")
    locked_fill = PatternFill("solid", fgColor="7C3AED")
    center = Alignment(horizontal="center", vertical="center")

    # Generate date range
    days = [week_start + timedelta(days=i) for i in range((week_end - week_start).days + 1)]
    day_names = ["Dom", "Lun", "Mar", "Mié", "Jue", "Vie", "Sáb"]

    # Header row 1 — Branch info
    ws.merge_cells(f"A1:{get_column_letter(3 + len(days))}1")
    ws["A1"] = f"{data['branch_name']} ({data['branch_code']}) | Semana {week_start} – {week_end}"
    ws["A1"].font = Font(bold=True, size=13)

    # Header row 2 — Columns
    headers = ["RUT", "Nombre", "Cargo"] + [f"{day_names[d.weekday() % 7 + 1 if d.weekday() < 6 else 0]}\n{d.day}/{d.month}" for d in days] + ["Total Hrs"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center

    # Data rows
    for row_i, worker in enumerate(data["workers"], 3):
        ws.cell(row=row_i, column=1, value=worker["rut"]).alignment = center
        ws.cell(row=row_i, column=2, value=worker["name"])
        ws.cell(row=row_i, column=3, value=worker.get("position", ""))
        for col_i, d in enumerate(days, 4):
            day_data = worker["days"].get(str(d))
            if day_data:
                if day_data["is_locked"] or day_data["is_day_off"]:
                    val = "🔒 Libre" if day_data["is_locked"] else "Libre"
                    ws.cell(row=row_i, column=col_i, value=val).fill = day_off_fill if day_data["is_day_off"] else locked_fill
                else:
                    code = day_data.get("shift_code", "")
                    hours = day_data.get("total_hours", 0)
                    ws.cell(row=row_i, column=col_i, value=f"{code} ({hours}h)").alignment = center
            else:
                ws.cell(row=row_i, column=col_i, value="—").alignment = center

        total_col = 4 + len(days)
        ws.cell(row=row_i, column=total_col, value=worker["total_hours_scheduled"]).alignment = center

    # Column widths
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 16
    for i in range(len(days)):
        ws.column_dimensions[get_column_letter(4 + i)].width = 11
    ws.column_dimensions[get_column_letter(4 + len(days))].width = 10

    # Stream the file
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"turnos_{data['branch_code']}_{week_start}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
